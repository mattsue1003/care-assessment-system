"""匯出：Word（python-docx）、PDF（Chrome 無頭列印，版面與預覽一致）、Excel（openpyxl）。"""
import io
import os
import shutil
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import schema as S

CHROME_CANDIDATES = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    "/Applications/Microsoft Edge Dev.app/Contents/MacOS/Microsoft Edge Dev",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
]
FONT = "微軟正黑體"


def _roc(s, sep="/"):
    if not s:
        return ""
    y, m, d = (int(x) for x in str(s)[:10].split("-"))
    return f"{y - 1911}{sep}{m:02d}{sep}{d:02d}"


def _roc_long(s):
    if not s:
        return "　　 年　 月　 日"
    y, m, d = (int(x) for x in str(s)[:10].split("-"))
    return f"{y - 1911} 年 {m} 月 {d} 日"


# ================================================================== PDF
def _chrome():
    for p in CHROME_CANDIDATES:
        if os.path.exists(p):
            return p
    p = shutil.which("google-chrome") or shutil.which("chromium")
    if not p:
        raise RuntimeError("找不到 Google Chrome / Edge，無法產生 PDF（可改用 Word 或「列印 → 另存 PDF」）")
    return p


def html_to_pdf(html):
    with tempfile.TemporaryDirectory() as td:
        src, out = os.path.join(td, "in.html"), os.path.join(td, "out.pdf")
        with open(src, "w", encoding="utf-8") as fh:
            fh.write(html)
        subprocess.run([_chrome(), "--headless=new", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
                        f"--user-data-dir={os.path.join(td, 'prof')}", "--no-pdf-header-footer",
                        "--print-to-pdf-no-header", f"--print-to-pdf={out}", "file://" + src],
                       check=True, capture_output=True, timeout=90)
        with open(out, "rb") as fh:
            return fh.read()


def html_to_pdf_many(jobs, workers=4):
    with ThreadPoolExecutor(workers) as ex:
        pdfs = list(ex.map(lambda j: html_to_pdf(j[1]), jobs))
    return [(j[0], p) for j, p in zip(jobs, pdfs)]


# ================================================================== Word
def _doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.2)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.15
    return doc


def _shade(cell, color="E8E8E8"):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color)
    tcPr.append(sh)


def _txt(cell, text, bold=False, size=None, center=False, color=None):
    cell.text = ""
    lines = str(text or "").split("\n")
    p = cell.paragraphs[0]
    for i, ln in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        r = p.add_run(ln)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return cell


def _head(doc, center, title, sub=None):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(center); r.font.size = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(18)
    if sub:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(sub); r.font.size = Pt(9)
    doc.add_paragraph()


def _table(doc, rows, cols, widths=None, style="Table Grid"):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = style
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def _kv(doc, pairs, ncol=2):
    """[(label, value), ...] → 2 組 label/value 一列"""
    rows = (len(pairs) + ncol - 1) // ncol
    t = _table(doc, rows, ncol * 2, [2.6, 6.4] * ncol)
    for i, (k, v) in enumerate(pairs):
        r, c = divmod(i, ncol)
        _shade(_txt(t.cell(r, c * 2), k, bold=True, center=True))
        _txt(t.cell(r, c * 2 + 1), v)
    return t


def _wide(doc, label, text):
    t = _table(doc, 1, 2, [2.6, 15.4])
    _shade(_txt(t.cell(0, 0), label, bold=True, center=True))
    _txt(t.cell(0, 1), text, size=10)


def _sign(doc, labels):
    doc.add_paragraph()
    p = doc.add_paragraph("　　".join(f"{l}：＿＿＿＿＿＿" for l in labels))
    p.runs[0].font.size = Pt(11)


def _save(doc):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def docx_d61(a, c, setting):
    doc = _doc()
    _head(doc, setting("center_name"), "職能治療評估表", "對應 care99 系統：D61.職能治療評估")
    _kv(doc, [("評估對象", c["label"]), ("評估日期", _roc_long(a["date"])),
              ("填表人", a["assessor"]), ("評估類別", a.get("kind") or ""),
              ("性別／年齡", f"{c.get('gender') or ''}／{c.get('age') or ''} 歲"), ("下次再評估", _roc(a.get("next_due")))])
    _wide(doc, "整體評估\n建議", a.get("summary"))
    _wide(doc, "備註", a.get("note"))
    doc.add_paragraph()
    rows = S.d61_rows(a["data"])
    t = _table(doc, len(rows) + 1, 4, [1.1, 2.2, 11.4, 3.3])
    for i, h in enumerate(["序號", "主題", "評估項目", "備註"]):
        _shade(_txt(t.cell(0, i), h, bold=True, center=True))
    for r, (no, title, lines, note) in enumerate(rows, 1):
        _txt(t.cell(r, 0), no, center=True)
        _txt(t.cell(r, 1), title, bold=True, center=True)
        _txt(t.cell(r, 2), "\n".join(lines), size=9.5)
        _txt(t.cell(r, 3), note, size=9)
    _sign(doc, ["職能治療師簽章", "主管覆核"])
    return _save(doc)


def docx_icope(a, c, setting):
    d = a["data"]
    j = S.icope_auto(d)
    doc = _doc()
    _head(doc, setting("center_name"), "職能治療 ICOPE 評估（長者六力）", "依國健署「長者功能評估」六大面向：認知、行動、營養、視力、聽力、憂鬱")
    _kv(doc, [("個案姓名", c["name"]), ("案號", c.get("case_no") or ""),
              ("性別／年齡", f"{c.get('gender') or ''}／{c.get('age') or ''} 歲"), ("出生日期", _roc(c.get("birth"))),
              ("評估日期", _roc_long(a["date"])), ("評估類別", a.get("kind") or ""),
              ("督導", f"{setting('supervisor')}（{setting('supervisor_title')}）"), ("評估者", a["assessor"])])
    doc.add_paragraph()
    rows = S.icope_rows(d)
    t = _table(doc, len(rows) + 1, 4, [2.0, 9.6, 1.6, 4.8])
    for i, h in enumerate(["面向", "評估內容與結果", "判定", "後續處置／建議"]):
        _shade(_txt(t.cell(0, i), h, bold=True, center=True))
    for r, (k, lbl, lines, finding, acts) in enumerate(rows, 1):
        _shade(_txt(t.cell(r, 0), lbl + ("\n（後續評估）" if k in ("med", "soc") else ""), bold=True, center=True), "F4F4F4")
        _txt(t.cell(r, 1), "\n".join(lines + ([f"說明：{finding}"] if finding else [])), size=9.5)
        jj = j.get(k, "")
        _txt(t.cell(r, 2), jj, bold=True, center=True, color="C0392B" if jj in ("異常", "需追蹤") else None)
        _txt(t.cell(r, 3), "\n".join("▸ " + x for x in acts), size=9.5)
    ab = [l for k, l in S.ICOPE_DOMAINS if j.get(k) == "異常"]
    ok = [l for k, l in S.ICOPE_DOMAINS if j.get(k) == "正常"]
    doc.add_paragraph()
    _wide(doc, "本次結果", f"異常 {len(ab)} 項（{'、'.join(ab) or '無'}）　正常 {len(ok)} 項（{'、'.join(ok) or '無'}）")
    _wide(doc, "目標", d.get("goal_text", ""))
    _wide(doc, "職能治療\n處置建議", d.get("plan_text", ""))
    _wide(doc, "資源連結", d.get("resource_text", ""))
    _wide(doc, "下次再評估", f"{_roc_long(a.get('next_due'))}（如跌倒、住院或功能明顯退化，提前再評估）")
    _sign(doc, ["評估者簽章", "個案／家屬", "主管覆核"])
    return _save(doc)


def docx_client(b, setting):
    c = b["c"]
    doc = _doc()
    _head(doc, setting("center_name"), "個案職能治療評估彙整", f"製表日期：{_roc_long(date.today().isoformat())}")
    _kv(doc, [("姓名", c["name"]), ("案號", c.get("case_no") or ""), ("性別／年齡", f"{c.get('gender') or ''}／{c.get('age') or ''} 歲"),
              ("出生日期", _roc(c.get("birth"))), ("現有服務", c.get("services") or ""), ("主責督導", c.get("supervisor") or ""),
              ("最近評估", _roc(b["due"]["last"])), ("下次應評估", f"{_roc(b['due']['next'])}（{b['due']['status']}）")])

    def h(text):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(12.5)

    h("一、歷次評估一覽")
    rows = [a for a in b["allas"]]
    t = _table(doc, len(rows) + 1, 4, [2.4, 3.2, 2.6, 9.8])
    for i, x in enumerate(["日期", "評估", "評估者", "摘要"]):
        _shade(_txt(t.cell(0, i), x, bold=True, center=True))
    names = {"D61": "D61 職能治療評估", "ICOPE": "六力 ICOPE", "D61_HIST": "D61（歷史）", "ICOPE_HIST": "六力（歷史）"}
    for r, a in enumerate(rows, 1):
        _txt(t.cell(r, 0), _roc(a["date"]), center=True)
        _txt(t.cell(r, 1), names.get(a["type"], a["type"]))
        _txt(t.cell(r, 2), a["assessor"] or "")
        summ = a["summary"] or a.get("note") or ""
        if a["type"] == "ICOPE":
            ab = S.icope_abnormal(a["data"])
            summ = f"異常：{'、'.join(ab) or '無'}"
        _txt(t.cell(r, 3), summ[:160] + ("…" if len(summ) > 160 else ""), size=9)

    tr = b["trend"]
    if tr["dates"]:
        h("二、六力歷次結果")
        t = _table(doc, len(tr["rows"]) + 2, len(tr["dates"]) + 1)
        _shade(_txt(t.cell(0, 0), "面向", bold=True, center=True))
        for i, d_ in enumerate(tr["dates"], 1):
            _shade(_txt(t.cell(0, i), _roc(d_), bold=True, center=True))
        for r, (lbl, vals) in enumerate(tr["rows"], 1):
            _txt(t.cell(r, 0), lbl, bold=True)
            for i, v in enumerate(vals, 1):
                _txt(t.cell(r, i), v or "—", center=True, color="C0392B" if v in ("異常", "需追蹤") else None)
        r = len(tr["rows"]) + 1
        _txt(t.cell(r, 0), "起立坐下（秒）", bold=True)
        for i, v in enumerate(tr["secs"], 1):
            _txt(t.cell(r, i), f"{v:g}" if v is not None else "—", center=True)

    dt = b["d61trend"]
    if dt["dates"]:
        h("三、D61 職能治療評估歷次比較")
        t = _table(doc, len(dt["rows"]) + 1, len(dt["dates"]) + 1)
        _shade(_txt(t.cell(0, 0), "項目", bold=True, center=True))
        for i, d_ in enumerate(dt["dates"], 1):
            _shade(_txt(t.cell(0, i), _roc(d_), bold=True, center=True))
        for r, (lbl, vals, opts) in enumerate(dt["rows"], 1):
            _txt(t.cell(r, 0), lbl, bold=True, size=9.5)
            for i, v in enumerate(vals, 1):
                bad = opts and v and v != opts[0]
                _txt(t.cell(r, i), v or "—", center=True, size=9.5, color="C0392B" if bad else None)
    if b["changes"]:
        p = doc.add_paragraph(); p.add_run("最近兩次 D61 變化：").bold = True
        for lbl, a1, a2, k in b["changes"]:
            doc.add_paragraph(f"・{lbl}：{a1} → {a2}（{k}）")

    h("四、輔具使用（評鑑 B9）")
    if b["aids"]:
        t = _table(doc, len(b["aids"]) + 1, 7)
        for i, x in enumerate(["輔具", "來源", "開始使用", "適配評估", "使用訓練", "最近檢查", "狀態"]):
            _shade(_txt(t.cell(0, i), x, bold=True, center=True, size=9.5))
        for r, x in enumerate(b["aids"], 1):
            for i, v in enumerate([x["item"], x["source"], _roc(x["start_date"]), _roc(x["fit_date"]), _roc(x["train_date"]),
                                   _roc(x["check_date"]), x["status"]]):
                _txt(t.cell(r, i), v or "—", size=9.5, center=i > 1)
    else:
        doc.add_paragraph("無使用輔具紀錄。")

    h("五、自立支援／復能目標（評鑑 B10）")
    if b["goals"]:
        for gl in b["goals"]:
            p = doc.add_paragraph()
            p.add_run(f"【{gl['domain']}】{gl['goal']}").bold = True
            doc.add_paragraph(f"方法：{gl['method'] or ''}　基準：{gl['baseline'] or ''}　開始：{_roc(gl['start_date'])}　"
                              f"檢討日：{_roc(gl['review_date'])}　狀態：{gl['status']}")
            for rv in gl["reviews"]:
                doc.add_paragraph(f"　・{_roc(rv['date'])} {rv['progress']}：{rv['note'] or ''}（{rv['reviewer'] or ''}）")
    else:
        doc.add_paragraph("尚未建立目標。")
    _sign(doc, ["職能治療師", "督導（護理師）", "主管"])
    return _save(doc)


# ================================================================== Excel
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFILL = PatternFill("solid", fgColor="1F6E5C")
FILLS = {"逾期": "F8D7DA", "30天內到期": "D6E9F8", "未評估": "FFF3CD", "異常": "F8D7DA", "需追蹤": "FDE8D4",
         "退步": "F8D7DA", "缺漏": "FFF3CD"}


def _sheet(wb, title, headers, rows, widths=None, note=None, color_cols=()):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(bold=True, size=12)
        r0 = 3
    for i, h in enumerate(headers, 1):
        c = ws.cell(r0, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HFILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for ri, row in enumerate(rows, r0 + 1):
        for ci, v in enumerate(row, 1):
            c = ws.cell(ri, ci, v)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if ci - 1 in color_cols and isinstance(v, str):
                for key, col in FILLS.items():
                    if key in v:
                        c.fill = PatternFill("solid", fgColor=col)
                        break
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 3)
    if rows:
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{r0}:{r0}"
    return ws


def xlsx_all(p, roc, age_at):
    wb = Workbook()
    wb.remove(wb.active)
    ind = p["ind"]
    period = f"{_roc(p['f']) if p['f'] > '1900' else '全部'} ～ {_roc(p['t']) if p['t'] < '2999' else '今'}"

    # 指標摘要
    ws = wb.create_sheet("指標摘要")
    rows = [
        (f"{p['center']}　評鑑佐證總表", ""),
        ("資料期間", period),
        ("製表日期", _roc(date.today().isoformat())),
        ("", ""),
        ("指標（115 年 附件2-3 小規模多機能）", "現況"),
        ("在案個案數", f"{ind['n']} 位"),
        ("B1 近 6 個月內完成評估／再評估", f"{ind['b1'][0]} 位（{ind['b1'][1]}%）"),
        ("B9 使用中輔具：開始日期＋適配評估＋使用訓練 皆有紀錄", f"{ind['b9'][0]}／{ind['b9'][1]} 件（{ind['b9'][2]}%）"),
        ("B10 有進行中自立支援目標的個案", f"{ind['b10'][0]} 位（{ind['b10'][1]}%）"),
        ("B10 近 3 個月內有檢討紀錄的目標", f"{ind['b10'][2]}／{ind['b10'][3]} 項"),
        ("", ""),
        ("期間內 D61 職能治療評估", f"{len(p['d61'])} 筆"),
        ("期間內 六力 ICOPE 評估", f"{len(p['ic'])} 筆"),
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    for c in ("A5", "B5"):
        ws[c].font = Font(bold=True, color="FFFFFF"); ws[c].fill = HFILL
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 30

    # 在案個案追蹤
    gl_count = {}
    for gl in p["goals"]:
        if gl["status"] == "進行中":
            gl_count[gl["client_id"]] = gl_count.get(gl["client_id"], 0) + 1
    aids_by = {}
    for a in p["aids"]:
        if a["status"] == "使用中":
            aids_by.setdefault(a["client_id"], []).append(a["item"])
    _sheet(wb, "在案個案評估追蹤",
           ["案號", "姓名", "性別", "年齡", "現有服務", "主責督導", "最近D61", "最近六力", "最近評估", "下次應評估", "評估狀態",
            "使用中輔具", "進行中目標", "備註"],
           [[c["case_no"], c["name"], c["gender"], c["age"], c["services"], c["supervisor"], roc(c["due"]["last_d61"]),
             roc(c["due"]["last_icope"]), roc(c["due"]["last"]), roc(c["due"]["next"]), c["due"]["status"],
             "、".join(aids_by.get(c["id"], [])), gl_count.get(c["id"], 0), c["note"]] for c in p["cs"]],
           [8, 10, 5, 5, 16, 9, 10, 10, 10, 10, 11, 18, 8, 24], color_cols=(10,))

    # D61
    keys = ["hobby_now", "adl_help", "func", "soc_coop", "soc_motiv"] + ["rom_" + k for k, _ in S.SIDES] + \
           ["mmt_" + k for k, _ in S.SIDES] + ["tone_ue", "tone_le", "sense_touch", "sense_hear", "sense_vision", "bpsd",
            "cog_alert", "cog_orient", "cog_attn", "cog_mem_s", "cog_mem_l", "cog_expr", "cog_comp", "cog_read", "cog_solve",
            "ot_goal", "ot_plan", "sched_am", "sched_pm", "sched_ev"]
    names = {"hobby_now": "興趣(現在)", "adl_help": "ADL需協助", "func": "功能性動作", "bpsd": "精神行為",
             "ot_goal": "治療目標", "ot_plan": "治療計畫"}
    heads = ["日期", "案號", "姓名", "填表人", "類別"] + [names.get(k) or S.D61_FIELDS[k]["qlabel"] for k in keys]
    heads += ["整體評估建議", "已登打care99"]
    rows = []
    for a in p["d61"]:
        b = S.d61_brief(a["data"], keys)
        rows.append([roc(a["date"]), a["case_no"], a["name"], a["assessor"], a["kind"]] + [b[k] for k in keys] +
                    [a["summary"], "是" if a["keyed"] else "否"])
    _sheet(wb, "D61職能治療評估", heads, rows, [10, 7, 9, 9, 7] + [12] * len(keys) + [60, 8])

    # 六力
    rows = []
    stat = {k: [0, 0] for k, _ in S.ICOPE_ALL}
    latest = {}
    for a in p["ic"]:
        latest[a["client_id"]] = a
        j = S.icope_auto(a["data"])
        ab = [l for k, l in S.ICOPE_DOMAINS if j.get(k) == "異常"]
        rows.append([roc(a["date"]), a["case_no"], a["name"], age_at(a["birth"], a["date"]), a["assessor"], a["kind"]] +
                    [j.get(k, "") for k, _ in S.ICOPE_ALL] + [a["data"].get("mob_seconds", ""), len(ab), "、".join(ab),
                     a["data"].get("plan_text", ""), roc(a["next_due"])])
    _sheet(wb, "六力ICOPE評估", ["日期", "案號", "姓名", "年齡", "評估者", "類別"] + [l for _, l in S.ICOPE_ALL] +
           ["起立坐下秒數", "異常項數", "異常面向", "處置建議", "下次再評估"], rows,
           [10, 7, 9, 5, 9, 7] + [8] * 8 + [8, 7, 18, 50, 10], color_cols=tuple(range(6, 14)))
    for a in latest.values():
        j = S.icope_auto(a["data"])
        for k, _ in S.ICOPE_ALL:
            if j.get(k):
                stat[k][1] += 1
                if j[k] in ("異常", "需追蹤"):
                    stat[k][0] += 1
    _sheet(wb, "六力統計", ["面向", "異常／需追蹤人數", "受評人數", "比例"],
           [[l, stat[k][0], stat[k][1], f"{round(stat[k][0] * 100 / stat[k][1])}%" if stat[k][1] else "—"]
            for k, l in S.ICOPE_ALL], [18, 16, 10, 10],
           note=f"六力統計（每位個案取期間內最近一次，共 {len(latest)} 位）")

    # 輔具
    rows = []
    for a in p["aids"]:
        miss = [n for k, n in (("start_date", "開始日期"), ("fit_date", "適配評估"), ("train_date", "使用訓練")) if not a[k]]
        rows.append([a["case_no"], a["name"], a["item"], a["category"], a["source"], roc(a["start_date"]), roc(a["fit_date"]),
                     roc(a["train_date"]), roc(a["check_date"]), a["status"], a["note"],
                     ("缺漏：" + "、".join(miss)) if miss and a["status"] == "使用中" else "完整"])
    _sheet(wb, "B9輔具清單", ["案號", "姓名", "輔具／用品", "類別", "來源", "開始使用", "適配評估日", "使用訓練日", "最近檢查日",
                            "狀態", "備註", "佐證完整性"], rows, [7, 9, 16, 8, 9, 10, 10, 10, 10, 8, 20, 22], color_cols=(11,))

    # 自立支援
    rows = []
    for gl in p["goals"]:
        last = gl["reviews"][-1] if gl["reviews"] else None
        rows.append([gl["case_no"], gl["name"], gl["domain"], gl["goal"], gl["method"], gl["baseline"], roc(gl["start_date"]),
                     roc(gl["review_date"]), gl["status"], roc(last["date"]) if last else "", last["progress"] if last else "",
                     "\n".join(f"{roc(r['date'])} {r['progress']}：{r['note'] or ''}" for r in gl["reviews"])])
    _sheet(wb, "B10自立支援目標", ["案號", "姓名", "面向", "目標", "方法", "基準（評估結果）", "開始日", "預定檢討", "狀態",
                              "最近檢討", "進展", "檢討紀錄"], rows, [7, 9, 12, 30, 28, 20, 10, 10, 8, 10, 10, 40])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
