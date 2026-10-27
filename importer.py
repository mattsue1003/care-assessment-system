"""從「福喜緣_在案個案名冊_職能治療追蹤_*.xlsx」匯入個案、系統 D61 歷史日期、B2c 輔具。可重複執行（不會重複建立）。"""
import glob
import os
from datetime import date, datetime

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))


def default_roster():
    files = sorted(glob.glob(os.path.join(BASE, "..", "00_個案名冊", "福喜緣_在案個案名冊_職能治療追蹤_*.xlsx")))
    return files[-1] if files else None


def _iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if not v:
        return ""
    s = str(v).strip().replace("-", "/")
    try:
        y, m, d = (int(x) for x in s.split("/")[:3])
        return date(y + 1911 if y < 1000 else y, m, d).isoformat()
    except ValueError:
        return ""


def _s(v):
    return "" if v is None else str(v).strip()


def run(con, path):
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    wb = openpyxl.load_workbook(path, data_only=True)
    added = updated = hist = aids = 0

    ws = wb["在案個案"]
    hdr_row = next(i for i, r in enumerate(ws.iter_rows(values_only=True), 1) if r and r[0] == "序號")
    hdr = [c.value for c in ws[hdr_row]]
    col = {h: i for i, h in enumerate(hdr) if h}
    for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        if not r or not r[col["姓名"]] or not isinstance(r[col["序號"]], (int, float)):
            continue  # 跳過底部統計列
        v = {
            "case_no": _s(r[col["案號"]]), "name": _s(r[col["姓名"]]), "gender": _s(r[col["性別"]]),
            "birth": _iso(r[col["生日(民國)"]]), "services": _s(r[col["現有服務"]]), "supervisor": _s(r[col["主責督導"]]),
            "living": _s(r[col["居住情況"]]), "identity": _s(r[col["身份別"]]) + ("／" + _s(r[col["收入別"]]) if r[col["收入別"]] else ""),
            "phone": _s(r[col["電話(主)"]]), "phone2": _s(r[col["電話(次)"]]), "address": _s(r[col["通訊地址"]]),
            "note": _s(r[col["OT備註"]]) if "OT備註" in col else "",
        }
        ex = con.execute("SELECT id FROM clients WHERE case_no=? AND name=?", (v["case_no"], v["name"])).fetchone()
        if ex:
            con.execute("""UPDATE clients SET gender=?, birth=?, services=?, supervisor=?, living=?, identity=?, phone=?,
                phone2=?, address=?, updated=? WHERE id=?""", (v["gender"], v["birth"], v["services"], v["supervisor"],
                v["living"], v["identity"], v["phone"], v["phone2"], v["address"], ts, ex[0]))
            updated += 1
        else:
            con.execute("""INSERT INTO clients (case_no,name,gender,birth,services,supervisor,living,identity,phone,phone2,
                address,note,status,created,updated) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,'在案',?,?)""",
                        (v["case_no"], v["name"], v["gender"], v["birth"], v["services"], v["supervisor"], v["living"],
                         v["identity"], v["phone"], v["phone2"], v["address"], v["note"], ts, ts))
            added += 1

    def cid_by_name(name):
        r = con.execute("SELECT id FROM clients WHERE name=?", (name,)).fetchone()
        return r[0] if r else None

    def add_hist(cid, typ, d, who, note):
        nonlocal hist
        if not cid or not d:
            return
        if con.execute("SELECT 1 FROM assessments WHERE client_id=? AND type IN (?,?) AND date=?",
                       (cid, typ, typ.replace("_HIST", ""), d)).fetchone():
            return
        con.execute("""INSERT INTO assessments (client_id,type,date,assessor,kind,data,summary,note,keyed,created,updated)
            VALUES (?,?,?,?,'歷史紀錄','{}','',?,1,?,?)""", (cid, typ, d, who, note, ts, ts))
        hist += 1

    if "D61系統紀錄" in wb.sheetnames:
        for r in wb["D61系統紀錄"].iter_rows(min_row=2, values_only=True):
            if r and r[2]:
                add_hist(cid_by_name(_s(r[2])), "D61_HIST", _iso(r[0]), _s(r[3]), "care99 系統 D61 紀錄（僅日期）")
    # 紙本六力日期
    ws = wb["在案個案"]
    for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        if r and r[col["姓名"]] and isinstance(r[col["序號"]], (int, float)):
            cid = cid_by_name(_s(r[col["姓名"]]))
            if "六力評估最近日期" in col:
                add_hist(cid, "ICOPE_HIST", _iso(r[col["六力評估最近日期"]]), "", "紙本六力評估（僅日期）")
            if "紙本再評估最近日期" in col:
                add_hist(cid, "D61_HIST", _iso(r[col["紙本再評估最近日期"]]), "", "紙本再評估（僅日期）")

    if "輔具B2c" in wb.sheetnames:
        for r in wb["輔具B2c"].iter_rows(min_row=2, values_only=True):
            if not r or not r[0] or not r[1]:
                continue
            cid = cid_by_name(_s(r[0]))
            if not cid or con.execute("SELECT 1 FROM aids WHERE client_id=? AND item=?", (cid, _s(r[1]))).fetchone():
                continue
            item = _s(r[1])
            cat = "行動" if any(k in item for k in ("杖", "助行", "輪椅", "爬梯")) else ("進食" if any(k in item for k in ("湯匙", "碗", "筷", "杯")) else "其他")
            con.execute("""INSERT INTO aids (client_id,item,category,source,start_date,status,note,created,updated)
                VALUES (?,?,?,?,?,'使用中',?,?,?)""", (cid, item, cat, _s(r[4]), _iso(r[3]), _s(r[2]), ts, ts))
            aids += 1
    return f"匯入完成：新增個案 {added} 位、更新 {updated} 位、歷史評估日期 {hist} 筆、輔具 {aids} 筆"
